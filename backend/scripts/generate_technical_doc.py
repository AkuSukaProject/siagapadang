import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def generate_technical_docx():
    doc = docx.Document()

    # Set Margins (1 inch / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Document Header Badge
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_badge = p_badge.add_run("DOKUMEN TEKNIS PENDUKUNG — GEMASTIK XIX 2026\nDIVISI VIII: PENGEMBANGAN PERANGKAT LUNAK")
    r_badge.font.size = Pt(9)
    r_badge.font.bold = True
    r_badge.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Panduan Instalasi, Penggunaan, dan Spesifikasi Teknis")
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Subtitle / Team Info
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run(
        "Karya: Siaga Padang — Sistem Navigasi Evakuasi Tsunami Luring Berbasis Pemodelan Pergerakan Kolektif\n"
        "Institusi: Universitas Andalas, Padang\n"
        "Tim Pengusul: Mikail Samyth Habibillah (Ketua) · Muhammad Habib (Anggota) · Sheva Ramadhan (Anggota)\n"
        "Dosen Pembimbing: Nisa Dwi Angresti, S.Si., M.Kom."
    )
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # Divider Line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    r_div = p_div.add_run("━" * 58)
    r_div.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

    # ==========================================
    # BAGIAN A: LATAR BELAKANG
    # ==========================================
    h_a = doc.add_heading("A. LATAR BELAKANG", level=1)
    h_a.paragraph_format.space_before = Pt(14)
    h_a.paragraph_format.space_after = Pt(6)
    h_a.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_a1 = doc.add_paragraph()
    p_a1.add_run("1.1 Kondisi Kebencanaan & Zona Rawan Kota Padang\n").bold = True
    p_a1.add_run(
        "Kota Padang berada di pesisir barat Pulau Sumatera dan berhadapan langsung dengan zona subduksi aktif Megathrust Mentawai–Siberut. "
        "BMKG mengidentifikasi segmen ini sebagai seismic gap dengan potensi pelepasan energi gempa bumi hingga skala M8,9. "
        "Dalam skenario terburuk, pemodelan BMKG mengestimasikan gelombang tsunami dapat mencapai pesisir Kota Padang dalam waktu kurang dari 30 menit "
        "dengan ketinggian limpasan melebihi 10 meter. Berdasarkan Kajian Risiko Bencana (KRB) Kota Padang Tahun 2023, ancaman rendaman tsunami "
        "mencakup 8 kecamatan dan 55 kelurahan dengan 242.750 jiwa terpapar (termasuk 77.014 kelompok rentan dan 637 penyandang disabilitas)."
    )

    p_a2 = doc.add_paragraph()
    p_a2.add_run("1.2 Urgensi Golden Time & Batasan Waktu Evakuasi\n").bold = True
    p_a2.add_run(
        "Materi sosialisasi resmi BPBD Kota Padang menetapkan rentang 20–30 menit sebagai golden time untuk melakukan evakuasi mandiri segera setelah guncangan kuat. "
        "Saat kondisi darurat pascagempa, masyarakat mengalami tekanan psikologis berat (cognitive freeze) dan kepanikan massal. "
        "Kondisi ini menuntut adanya sistem panduan yang mampu memberikan arahan arah evakuasi seketika tanpa memerlukan analisis orientasi peta yang rumit."
    )

    p_a3 = doc.add_paragraph()
    p_a3.add_run("1.3 Keterbatasan Solusi Eksisting (Peta Statis & Sistem Peringatan Daring)\n").bold = True
    p_a3.add_run(
        "1. Peta Evakuasi Statis (Padang Sigap): Disajikan dalam bentuk gambar raster per kelurahan tempat tinggal, sehingga tidak relevan jika warga sedang beraktivitas di luar domisili serta menuntut orientasi manual yang sulit saat panik.\n"
        "2. Ketergantungan Internet (InaTEWS BMKG & Google EEW): Pada gempa berskala besar, menara BTS seluler dan pasokan listrik hampir dipastikan padam total (single point of failure). Peringatan dini juga hanya memberi tahu bahwa gempa terjadi, bukan ke mana warga harus melangkah.\n"
        "3. Bottleneck Rute Individual: Algoritma penentuan rute konvensional (Dijkstra terpendek) mengarahkan massa ke 1 Tempat Evakuasi Sementara (TES) terdekat secara serentak, memicu overcapacity fatal sementara gedung TES lain masih kosong."
    )

    # ==========================================
    # BAGIAN B: TUJUAN
    # ==========================================
    h_b = doc.add_heading("B. TUJUAN DAN MANFAAT", level=1)
    h_b.paragraph_format.space_before = Pt(14)
    h_b.paragraph_format.space_after = Pt(6)
    h_b.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_b1 = doc.add_paragraph()
    p_b1.add_run("2.1 Tujuan Sistem\n").bold = True
    p_b1.add_run(
        "1. Mengembangkan sistem navigasi evakuasi tsunami berbasis mobile Android yang berfungsi 100% luring (offline-first) tanpa ketergantungan internet.\n"
        "2. Mengimplementasikan algoritma routing berbobot kapasitas (capacity-aware spatial routing) hasil prakomputasi multi-agen untuk mendistribusikan massa secara proporsional ke 143 TES resmi.\n"
        "3. Menyediakan navigasi 3 peringkat rute alternatif (Rank 1, 2, 3) yang dapat dialihkan secara instan saat terjadi hambatan jalan di lapangan.\n"
        "4. Merancang antarmuka darurat dengan beban kognitif minimal (standar aksesibilitas WCAG 2.1 AA kontras tinggi) serta widget layar utama 1-ketukan."
    )

    p_b2 = doc.add_paragraph()
    p_b2.add_run("2.2 Manfaat bagi Pengguna dan Stakeholder\n").bold = True
    p_b2.add_run(
        "• Bagi Masyarakat & Pendatang: Memperoleh panduan evakuasi langkah demi langkah dalam golden time 20–30 menit dari koordinat keberadaan nyata tanpa risiko kehilangan arah.\n"
        "• Bagi Pengelola Bencana (BPBD Kota Padang): Memperoleh analisis persebaran beban titik evakuasi sebelum bencana dan menerima laporan hambatan jalur secara crowd-sourced setelah jaringan pulih."
    )

    # ==========================================
    # BAGIAN C: NILAI INOVASI DAN DAMPAK
    # ==========================================
    h_c = doc.add_heading("C. NILAI INOVASI DAN DAMPAK PEMANFAATAN", level=1)
    h_c.paragraph_format.space_before = Pt(14)
    h_c.paragraph_format.space_after = Pt(6)
    h_c.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_c1 = doc.add_paragraph()
    p_c1.add_run("3.1 Pilar Inovasi Perangkat Lunak\n").bold = True
    p_c1.add_run(
        "1. Precomputed Capacity-Aware Routing: Komputasi rute dan simulasi 2.000 agen dilakukan di tahap persiapan (Reverse Single-Source Dijkstra + batasan kuota TES), sehingga ponsel pengguna hanya membaca hasil simpanan lokal tanpa menguras baterai dan CPU saat darurat.\n"
        "2. Pure Offline Vector Map (Zero-Network Cost): Menggunakan MapLibre Native dan SQLite lokal berukuran 70,18 MB untuk seluruh 8 kecamatan pesisir tanpa ketergantungan API eksternal berbayar.\n"
        "3. Multi-Tier Fallback Resilience: Mekanisme pengalihan rute otomatis dan berjenjang (Rank 1 → Rank 2 → Rank 3 → Azimuth Menjauhi Pantai) saat jalan tertutup puing.\n"
        "4. Safety-First Design: Peniadaan fitur pelacakan lokasi antar-keluarga secara sengaja guna mencegah kepanikan bergerak berlawanan arah kembali ke pesisir."
    )

    # Tabel Perbandingan Solusi Terdahulu
    p_tab1 = doc.add_paragraph()
    p_tab1.add_run("Tabel C.1 — Perbandingan dengan Sistem & Solusi Terdahulu").bold = True
    p_tab1.paragraph_format.space_before = Pt(6)
    p_tab1.paragraph_format.space_after = Pt(4)

    t_comp = doc.add_table(rows=1, cols=3)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t_comp.rows[0].cells
    hdr[0].text = "Sistem / Solusi Terdahulu"
    hdr[1].text = "Pendekatan & Keterbatasan"
    hdr[2].text = "Kebaruan & Keunggulan Siaga Padang"
    for i in range(3):
        hdr[i].paragraphs[0].runs[0].font.bold = True
        hdr[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(hdr[i], "0F172A")
        set_cell_margins(hdr[i], top=100, bottom=100, left=100, right=100)

    comp_data = [
        ("Peta Raster Padang Sigap", "Format gambar statis per kelurahan; menuntut orientasi manual.", "Interaktif berbasis koordinat GPS nyata; menyajikan rute langkah demi langkah."),
        ("InaRISK Personal (BNPB)", "Informasi tingkat bahaya umum; membutuhkan internet aktif.", "Sistem navigasi evakuasi personal; beroperasi 100% luring (offline-first)."),
        ("InaTEWS BMKG / Google EEW", "Peringatan dini gempa; bergantung koneksi seluler/BTS.", "Navigasi penunjuk arah mandiri; data BMKG dikonsumsi sebagai pelengkap saat online."),
        ("Navigasi Terpendek Konvensional", "Rute terpendek individu; memicu penumpukan overcapacity di 1 TES.", "Prakomputasi berbatas kapasitas; membagi beban ke 3 peringkat rute aman.")
    ]
    for row_data in comp_data:
        r = t_comp.add_row()
        r.cells[0].text = row_data[0]
        r.cells[0].paragraphs[0].runs[0].font.bold = True
        r.cells[1].text = row_data[1]
        r.cells[2].text = row_data[2]
        for idx in range(3):
            set_cell_background(r.cells[idx], "F8FAFC")
            set_cell_margins(r.cells[idx], top=80, bottom=80, left=100, right=100)

    # Tabel Dampak Kuantitatif
    p_tab2 = doc.add_paragraph()
    p_tab2.add_run("\nTabel C.2 — Dampak Kuantitatif Pemodelan Berbatas Kapasitas").bold = True
    p_tab2.paragraph_format.space_before = Pt(8)
    p_tab2.paragraph_format.space_after = Pt(4)

    t_dampak = doc.add_table(rows=1, cols=4)
    t_dampak.alignment = WD_TABLE_ALIGNMENT.CENTER
    d_hdr = t_dampak.rows[0].cells
    d_hdr_titles = ["Indikator Pengujian", "Pendekatan Terpendek", "Pendekatan Berbatas Kapasitas", "Selisih Dampak Nyata"]
    for i, title in enumerate(d_hdr_titles):
        d_hdr[i].text = title
        d_hdr[i].paragraphs[0].runs[0].font.bold = True
        d_hdr[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(d_hdr[i], "0284C7")
        set_cell_margins(d_hdr[i], top=100, bottom=100, left=100, right=100)

    dampak_rows = [
        ("Persentase Warga Tertampung (20 Menit)", "39,52%", "74,49%", "+34,97% Peningkatan Kapasitas"),
        ("Jumlah Titik TES Overcapacity", "52 titik TES", "0 titik TES", "-52 Titik Penumpukan Beban"),
        ("Waktu Respon Penyajian Jalur Luring", "> 3 detik (di HP)", "< 0,4 detik", "Seketika (< 1 Detik)"),
        ("Ukuran Database Spasial di HP", "-", "70,18 MB", "Sangat Efisien (Target ≤75 MB)"),
        ("Keberhasilan Unit Testing Logika", "-", "38/38 Kasus Lulus", "100% Lulus Pengujian")
    ]
    for d_row in dampak_rows:
        r = t_dampak.add_row()
        r.cells[0].text = d_row[0]
        r.cells[0].paragraphs[0].runs[0].font.bold = True
        r.cells[1].text = d_row[1]
        r.cells[2].text = d_row[2]
        r.cells[3].text = d_row[3]
        for idx in range(4):
            set_cell_background(r.cells[idx], "F8FAFC")
            set_cell_margins(r.cells[idx], top=80, bottom=80, left=100, right=100)

    # ==========================================
    # BAGIAN D: DESKRIPSI FUNGSIONAL & PANDUAN
    # ==========================================
    h_d = doc.add_heading("D. DESKRIPSI FUNGSIONAL, DETAIL FITUR, DAN PANDUAN PENGGUNAAN", level=1)
    h_d.paragraph_format.space_before = Pt(14)
    h_d.paragraph_format.space_after = Pt(6)
    h_d.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_d_spec = doc.add_paragraph()
    p_d_spec.add_run("4.1 Spesifikasi Kebutuhan Fungsional (F-01 s.d. F-14)").bold = True
    p_d_spec.paragraph_format.space_after = Pt(4)

    t_feat = doc.add_table(rows=1, cols=4)
    t_feat.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_hdr = t_feat.rows[0].cells
    f_hdr_titles = ["Kode", "Nama Fitur Fungsional", "Deskripsi Implementasi Teknis", "Mode Operasi"]
    for i, title in enumerate(f_hdr_titles):
        f_hdr[i].text = title
        f_hdr[i].paragraphs[0].runs[0].font.bold = True
        f_hdr[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(f_hdr[i], "0F172A")
        set_cell_margins(f_hdr[i], top=100, bottom=100, left=100, right=100)

    features = [
        ("F-01", "Widget Layar Utama (Home Screen)", "Menampilkan status zona & akses instan 1-ketukan ke navigasi darurat", "Offline Total"),
        ("F-02", "Penentuan Status Zona Pengguna", "Membaca koordinat GPS hardware dan mencocokkan dengan poligon rendaman", "Offline Total"),
        ("F-03", "Penyajian Jalur Evakuasi Luring", "Mengambil rute simpul terdekat pada Room SQLite lokal (< 1 detik)", "Offline Total"),
        ("F-04", "Penunjuk Arah Berbasis Kompas", "Menggunakan sensor magnetometer untuk memutar orientasi peta", "Offline Total"),
        ("F-05", "Penghitung Waktu Simulatif", "Menghitung mundur waktu evakuasi mandiri sejak aplikasi dibuka", "Offline Total"),
        ("F-06", "Deteksi Kedatangan di Titik Aman", "Konfirmasi kedatangan otomatis saat jarak < 20 meter pada 3 pembacaan", "Offline Total"),
        ("F-07", "Titik Temu Keluarga (Masa Tenang)", "Perencanaan titik kumpul keluarga saat kondisi normal/sebelum gempa", "Offline Total"),
        ("F-08", "Alih Jalur Alternatif Cepat", "Tombol lapor jalur terhalang untuk beralih instan ke rute Rank 2/Rank 3", "Offline Total"),
        ("F-09", "Status Terkini BMKG TEWS", "Mengambil parameter gempa resmi BMKG melalui API saat online", "Daring (Sync)"),
        ("F-10", "Konfirmasi Keselamatan (Check-in)", "Mengirimkan status selamat di TES ke server pengelola bencana", "Daring (Sync)"),
        ("F-11", "Pengiriman Laporan Hambatan Jalan", "Sinkronisasi koordinat hambatan jalan ke server saat jaringan pulih", "Daring (Sync)"),
        ("F-12", "Sinkronisasi Pembaruan Wilayah", "Mengunduh patch pembaruan data spasial berkala di masa tenang", "Daring (Sync)"),
        ("F-13", "Simulasi Pergerakan Kolektif", "Menghitung pergerakan multi-agen untuk identifikasi titik rawan padat", "Prakomputasi"),
        ("F-14", "Analisis Distribusi Kapasitas TES", "Menghitung rasio beban penduduk terhadap daya tampung gedung TES", "Prakomputasi")
    ]
    for f_item in features:
        r = t_feat.add_row()
        r.cells[0].text = f_item[0]
        r.cells[0].paragraphs[0].runs[0].font.bold = True
        r.cells[1].text = f_item[1]
        r.cells[1].paragraphs[0].runs[0].font.bold = True
        r.cells[2].text = f_item[2]
        r.cells[3].text = f_item[3]
        for idx in range(4):
            set_cell_background(r.cells[idx], "F8FAFC")
            set_cell_margins(r.cells[idx], top=80, bottom=80, left=100, right=100)

    # Panduan Instalasi
    p_inst = doc.add_paragraph()
    p_inst.add_run("\n4.2 Panduan Instalasi Perangkat Lunak\n").bold = True
    p_inst.add_run(
        "A. Kebutuhan Sistem Minimum:\n"
        "• Smartphone Android: Android 6.0+ (API 23), RAM 2 GB, Ruang Penyimpanan Kosong ≥250 MB, Sensor GPS & Magnetometer.\n"
        "• Server Peladen (Opsional/Testing): Python 3.11+, PostgreSQL 15+ dengan PostGIS 3.3+, RAM 4 GB.\n\n"
        "B. Langkah Pemasangan Aplikasi Android:\n"
        "1. Salin berkas 'siagapadang-release.apk' ke penyimpanan ponsel Android.\n"
        "2. Buka berkas APK dan pilih 'Install' (Izinkan 'Install from Unknown Sources' bila diminta).\n"
        "3. Buka aplikasi Siaga Padang dan berikan izin 'Lokasi Presisi' saat aplikasi pertama kali dijalankan.\n"
        "4. Basis data spasial Kota Padang (70,18 MB) telah terpasang secara otomatis di dalam aplikasi.\n"
        "5. Tambahkan Widget Siaga Padang ke Layar Utama (Home Screen) untuk akses darurat tercepat."
    )

    # Panduan Penggunaan
    p_use = doc.add_paragraph()
    p_use.add_run("\n4.3 Panduan Penggunaan dalam Berbagai Skenario\n").bold = True
    p_use.add_run(
        "• Skenario Tanggap Darurat Pascagempa:\n"
        "  1. Setelah guncangan kuat berhenti, segera keluar dari gedung ke ruang terbuka.\n"
        "  2. Buka aplikasi melalui Widget Layar Utama dengan 1 ketukan instan.\n"
        "  3. Aplikasi langsung mengunci koordinat GPS dan menyajikan rute evakuasi Rank 1 menuju TES terdekat (< 1 detik).\n"
        "  4. Ikuti arah panah panduan dan petunjuk belokan dengan berjalan cepat (jangan berlari dan jangan memakai kendaraan).\n"
        "• Skenario Jalur Terhalang Puing / Roboh:\n"
        "  1. Tekan tombol kuning 'Jalur Terhalang' pada antarmuka navigasi.\n"
        "  2. Sistem seketika mengalihkan panduan ke Rute Alternatif Rank 2 (atau Rank 3) tanpa proses komputasi ulang.\n"
        "• Skenario Tiba di Lokasi Aman (TES/TEA):\n"
        "  1. Saat berada dalam radius 20 meter dari gedung TES, aplikasi otomatis menampilkan Layar Kedatangan Aman.\n"
        "  2. Tetap berada di lantai atas gedung TES dan tunggu instruksi resmi dari petugas BPBD setempat."
    )

    # ==========================================
    # BAGIAN E: SCREENSHOT MOCKUP
    # ==========================================
    h_e = doc.add_heading("E. TANGKAPAN LAYAR PERANGKAT LUNAK (SCREENSHOTS)", level=1)
    h_e.paragraph_format.space_before = Pt(14)
    h_e.paragraph_format.space_after = Pt(6)
    h_e.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_e1 = doc.add_paragraph()
    p_e1.add_run(
        "1. Gambar E.1 — Widget Layar Utama Android (Home Screen Widget):\n"
        "   Menampilkan status zona pengguna (Aman / Waspada Rendaman) dan tombol darurat navigasi 1-ketukan.\n\n"
        "2. Gambar E.2 — Antarmuka Navigasi Evakuasi Luring (Active Navigation UI):\n"
        "   Menampilkan instruksi belokan berukuran besar, peta vektor luring rute ungu, estimasi waktu jalan kaki, dan tombol 'Jalur Terhalang'.\n\n"
        "3. Gambar E.3 — Layar Konfirmasi Kedatangan Aman (Safe Arrival Dialog):\n"
        "   Pemberitahuan otomatis saat pengguna tiba di TES dengan instruksi untuk tetap berada di lokasi vertikal.\n\n"
        "4. Gambar E.4 — Dokumentasi API Backend Peladen (FastAPI Swagger Docs):\n"
        "   Antarmuka interaktif endpoint REST API peladen di 'http://localhost:8000/docs' untuk sinkronisasi dataset, laporan hambatan, dan proxy BMKG."
    )

    output_path = "Dokumen_Teknis_Siaga_Padang_GEMASTIK_2026.docx"
    doc.save(output_path)
    print(f"[OK] Successfully generated: {output_path}")

if __name__ == "__main__":
    generate_technical_docx()
