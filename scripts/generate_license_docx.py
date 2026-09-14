import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Set Margins (1 inch / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Document Header / Badge
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_badge = p_badge.add_run("DOKUMEN PENDUKUNG — GEMASTIK XIX 2026")
    run_badge.font.size = Pt(9)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("Daftar Perangkat Lunak, Library, dan Lisensi Pihak Ketiga")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Document Subtitle / Metadata
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("Judul Karya: Siaga Padang: Sistem Navigasi Evakuasi Tsunami Luring Berbasis Pemodelan Pergerakan Kolektif\nFormat Deliverable: Nama | Versi | Fungsi/Peran | Lisensi")
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # Intro Paragraph
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(12)
    p_intro.add_run("Berikut adalah daftar lengkap komponen pihak ketiga (third-party libraries, frameworks, dan database engines) yang digunakan dalam pengembangan purwarupa ")
    r_bold = p_intro.add_run("Siaga Padang")
    r_bold.bold = True
    p_intro.add_run(". Seluruh komponen yang dipilih dipastikan bebas dari konflik lisensi berbayar untuk mematuhi persyaratan kompetisi GEMASTIK XIX 2026.")

    # Table Creation
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Column Widths
    col_widths = [Inches(0.4), Inches(1.8), Inches(1.0), Inches(2.3), Inches(1.0)]
    
    # Format Header Row
    hdr_cells = table.rows[0].cells
    hdr_titles = ["No", "Nama Komponen", "Versi", "Fungsi / Peran dalam Sistem", "Lisensi"]
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(hdr_cells[i], "0F172A") # Dark Slate
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
        hdr_cells[i].width = col_widths[i]

    # Data Items Grouped by Category
    data_groups = [
        {
            "category": "📱 CATEGORY 1: ANDROID FRONTEND (App Evakuasi Warga — gradle/libs.versions.toml)",
            "color": "0284C7", # Blue
            "items": [
                ("Kotlin", "2.3.21", "Bahasa pemrograman aplikasi Android", "Apache-2.0"),
                ("Jetpack Compose UI", "1.10.4 (Compose BOM 2026.02.01)", "Antarmuka pengguna deklaratif", "Apache-2.0"),
                ("Material 3", "1.4.0 (Compose BOM 2026.02.01)", "Komponen antarmuka Material Design", "Apache-2.0"),
                ("AndroidX Activity Compose", "1.12.4", "Integrasi Activity dengan Jetpack Compose", "Apache-2.0"),
                ("AndroidX Lifecycle Runtime Compose", "2.10.0", "Pengelolaan siklus hidup dan state antarmuka", "Apache-2.0"),
                ("AndroidX Lifecycle ViewModel Compose", "2.10.0", "Integrasi ViewModel dengan Compose", "Apache-2.0"),
                ("Room Runtime, KTX, & Compiler", "2.8.4", "Akses basis data SQLite lokal", "Apache-2.0"),
                ("MapLibre Native Android", "13.4.1", "Penyajian peta, jalur, zona, dan penanda", "BSD-2-Clause"),
                ("KotlinX Coroutines Android", "1.10.2", "Pemrosesan asinkron dan aliran data", "Apache-2.0"),
                ("Google Play Services Location", "21.4.0", "Pembacaan posisi melalui Fused Location Provider", "Google Mobile Dev Terms"),
                ("JUnit", "4.13.2", "Pengujian unit aplikasi Android", "EPL 1.0"),
            ]
        },
        {
            "category": "⚙️ CATEGORY 2: BACKEND SERVER & REST API SERVICES (requirements.txt)",
            "color": "059669", # Green
            "items": [
                ("FastAPI", "0.104.1", "Framework REST API Python berkinerja tinggi (ASGI)", "MIT"),
                ("Uvicorn", "0.24.0.post1", "Web server ASGI berkecepatan tinggi melayani request API Android", "BSD 3-Clause"),
                ("PostgreSQL", "15", "Engine basis data relasional utama peladen", "PostgreSQL License"),
                ("PostGIS", "3.4", "Ekstensi pengolah data spasial di level DB (ST_DWithin, ST_Covers)", "GPLv2"),
                ("SQLAlchemy", ">=2.0.35", "ORM untuk menghubungkan Python dan PostgreSQL", "MIT"),
                ("GeoAlchemy2", "0.14.2", "Ekstensi spasial SQLAlchemy untuk kueri geometri langsung dari Python", "MIT"),
                ("Psycopg 3", ">=3.2.0", "Adaptor PostgreSQL untuk Python berbasis asinkron", "LGPL-3.0"),
                ("Pydantic Settings", "2.1.0", "Pengelolaan konfigurasi aplikasi dan variabel lingkungan", "MIT"),
                ("Alembic", ">=1.12.1", "Alat kontrol versi dan migrasi skema basis data PostgreSQL", "MIT"),
                ("Pytest", ">=7.4.3", "Framework pengujian unit dan integrasi otomatis", "MIT"),
                ("Python-dotenv", "1.0.0", "Memuat variabel lingkungan dari file .env", "BSD 3-Clause"),
            ]
        },
        {
            "category": "🗺️ CATEGORY 3: SPATIAL COMPUTING & GRAPH ENGINE",
            "color": "7C3AED", # Purple
            "items": [
                ("OSMnx", "Tidak dikunci dalam requirements.txt", "Mengunduh, memodelkan, dan menganalisis jaringan jalan OpenStreetMap", "MIT"),
                ("NetworkX", "Tidak dikunci dalam requirements.txt", "Mesin komputasi graf matematis (Dijkstra algorithm) rute evakuasi", "BSD 3-Clause"),
                ("GeoPandas", "Tidak dikunci dalam requirements.txt", "Geospasial dataframes untuk membaca GeoJSON/Shapefile zona bahaya", "BSD 3-Clause"),
                ("Matplotlib", "Tidak dikunci dalam requirements.txt", "Visualisasi graf spasial dan peta evakuasi untuk laporan analisis", "PSF-based License"),
                ("Shapely", "Dependensi transitif; versi tidak dikunci", "Operasi geometri spasial 2D (intersection, polygon covers, distance)", "BSD 3-Clause"),
                ("pyproj", "Dependensi transitif; versi tidak dikunci", "Transformasi sistem koordinat Kartografi (EPSG:4326 ke EPSG:3857)", "MIT"),
                ("Fiona", "Dependensi transitif; versi tidak dikunci", "Pustaka I/O data geospasial untuk membaca berkas peta GIS", "BSD 3-Clause"),
            ]
        }
    ]

    row_count = 1
    for group in data_groups:
        # Category Header Row
        cat_row = table.add_row()
        cat_cell = cat_row.cells[0]
        cat_cell.merge(cat_row.cells[4])
        cat_cell.text = group["category"]
        p_cat = cat_cell.paragraphs[0]
        p_cat.runs[0].font.bold = True
        p_cat.runs[0].font.size = Pt(9.5)
        p_cat.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(cat_cell, group["color"])
        set_cell_margins(cat_cell, top=100, bottom=100, left=100, right=100)

        # Items
        for idx, item in enumerate(group["items"]):
            r = table.add_row()
            cells = r.cells
            
            # Row shading (alternate background)
            bg_color = "F8FAFC" if idx % 2 == 1 else "FFFFFF"

            data_tuple = (str(row_count), item[0], item[1], item[2], item[3])
            for c_i, val in enumerate(data_tuple):
                cells[c_i].text = val
                p = cells[c_i].paragraphs[0]
                p.runs[0].font.size = Pt(9)
                
                # Column specific styling
                if c_i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.runs[0].font.bold = True
                elif c_i == 1:
                    p.runs[0].font.bold = True
                    p.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                elif c_i == 2:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.runs[0].font.color.rgb = RGBColor(0x47, 0x55, 0x69)
                elif c_i == 4:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.runs[0].font.bold = True
                    if "MIT" in val or "Apache" in val or "BSD" in val:
                        p.runs[0].font.color.rgb = RGBColor(0x05, 0x96, 0x69)
                    else:
                        p.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

                set_cell_background(cells[c_i], bg_color)
                set_cell_margins(cells[c_i], top=80, bottom=80, left=100, right=100)
                cells[c_i].width = col_widths[c_i]

            row_count += 1

    # Conclusion & License Assurance Paragraph (Exact Text requested by Mikail)
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(16)
    p_foot.paragraph_format.space_after = Pt(6)
    r_foot_title = p_foot.add_run("Pernyataan Kepatuhan Lisensi Perangkat Lunak:\n")
    r_foot_title.bold = True
    r_foot_title.font.size = Pt(10)
    
    r_foot_body = p_foot.add_run(
        "Sebagian besar komponen pihak ketiga yang digunakan Siaga Padang merupakan perangkat lunak sumber terbuka dengan lisensi Apache-2.0, MIT, BSD, PostgreSQL License, LGPL, atau GPLv2. "
        "Google Play Services Location digunakan berdasarkan Google Mobile Developer Services Terms. "
        "Tidak terdapat pustaka komersial berbayar yang memerlukan biaya lisensi atau langganan untuk menjalankan fungsi inti aplikasi. "
        "Penggunaan setiap komponen tetap mengikuti ketentuan atribusi dan distribusi dari pemegang lisensinya."
    )
    r_foot_body.font.size = Pt(9.5)
    r_foot_body.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Save document
    filename = "Daftar_Library_dan_Lisensi_GEMASTIK_2026.docx"
    doc.save(filename)
    print(f"Document successfully created: {filename}")

if __name__ == "__main__":
    create_document()
